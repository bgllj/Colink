from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Literal

import openpyxl
import xlrd

from class_table_backend.domain.issues import Issue, IssueCode

_XLS_MAGIC = b"\xd0\xcf\x11\xe0"
_XLSX_MAGIC = b"PK\x03\x04"

# (rlo, rhi, clo, chi) 0-based, end-exclusive — same shape as xlrd's merged_cells.
MergeRange = tuple[int, int, int, int]


class WorkbookReadError(ValueError):
    """Raised when workbook bytes cannot be read; carries a structured Issue."""

    def __init__(self, issue: Issue):
        self.issue = issue
        super().__init__(issue.message)


@dataclass
class WorkbookData:
    sheets: dict[str, list[list[str]]] = field(default_factory=dict)
    merged: dict[str, list[MergeRange]] = field(default_factory=dict)


def detect_format(data: bytes) -> Literal["xls", "xlsx"] | None:
    if data[:4] == _XLS_MAGIC:
        return "xls"
    if data[:4] == _XLSX_MAGIC:
        return "xlsx"
    return None


def a1_coordinate(row: int, col: int) -> str:
    """Convert 0-based row/col to an A1-style coordinate."""
    n = col + 1
    letters = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        letters = chr(ord("A") + rem) + letters
    return f"{letters}{row + 1}"


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0 and value.microsecond == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value)


def _read_xls(data: bytes) -> WorkbookData:
    try:
        book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    except Exception as exc:  # noqa: BLE001 - surface any parser failure as structured issue
        raise WorkbookReadError(
            Issue(
                code=IssueCode.MALFORMED_WORKBOOK,
                message=f"无法解析的 xls 工作簿: {exc}",
            )
        ) from exc

    workbook = WorkbookData()
    for sheet in book.sheets():
        grid: list[list[str]] = []
        for row_index in range(sheet.nrows):
            row: list[str] = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                if cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append("")
                elif cell.ctype == xlrd.XL_CELL_BLANK:
                    row.append("")
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        as_datetime = xlrd.xldate_as_datetime(cell.value, book.datemode)
                    except Exception:  # noqa: BLE001
                        row.append(_stringify(cell.value))
                    else:
                        row.append(_stringify(as_datetime))
                else:
                    row.append(_stringify(cell.value))
            grid.append(row)
        workbook.sheets[sheet.name] = grid
        workbook.merged[sheet.name] = [
            (rlo, rhi, clo, chi) for (rlo, rhi, clo, chi) in sheet.merged_cells
        ]
    return workbook


def _read_xlsx(data: bytes) -> WorkbookData:
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001 - surface any parser failure as structured issue
        raise WorkbookReadError(
            Issue(
                code=IssueCode.MALFORMED_WORKBOOK,
                message=f"无法解析的 xlsx 工作簿: {exc}",
            )
        ) from exc

    workbook = WorkbookData()
    try:
        for sheet in book.worksheets:
            grid: list[list[str]] = []
            for row_index in range(1, (sheet.max_row or 0) + 1):
                row: list[str] = []
                for col_index in range(1, (sheet.max_column or 0) + 1):
                    row.append(_stringify(sheet.cell(row_index, col_index).value))
                grid.append(row)
            workbook.sheets[sheet.title] = grid
            workbook.merged[sheet.title] = [
                (merged.min_row - 1, merged.max_row, merged.min_col - 1, merged.max_col)
                for merged in sheet.merged_cells.ranges
            ]
    finally:
        book.close()
    return workbook


def read_workbook(data: bytes) -> WorkbookData:
    if not data:
        raise WorkbookReadError(
            Issue(code=IssueCode.MALFORMED_WORKBOOK, message="工作簿内容为空")
        )
    file_format = detect_format(data)
    if file_format == "xls":
        return _read_xls(data)
    if file_format == "xlsx":
        return _read_xlsx(data)
    raise WorkbookReadError(
        Issue(
            code=IssueCode.UNSUPPORTED_FILE_TYPE,
            message="无法识别的工作簿文件类型 (需要 OLE2 .xls 或 zip .xlsx)",
        )
    )
