from __future__ import annotations

import io

import openpyxl
import pytest
from sqlalchemy.orm import Session

from class_table_backend.import_flow.service import ImportService
from class_table_backend.import_flow.status import ImportStatus
from class_table_backend.persistence.db import get_engine, get_session_factory
from class_table_backend.persistence.tables import (
    Base,
    CourseRow,
    MeetingOccurrenceRow,
)


@pytest.fixture
def session() -> Session:
    engine = get_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = get_session_factory(engine)
    session = factory()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def _xlsx_with_courses(*cell_texts: str) -> bytes:
    """Build a minimal chengdu_wenli_v1 grid: title + weekday header + slot rows.

    Each arg is the 星期一 cell text for that slot row (may contain newlines).
    """
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.cell(1, 1, "成都文理学院2026-2027学年第一学期课表")
    sheet.cell(2, 1, "2026-08-31")
    sheet.cell(3, 1, "大节")
    sheet.cell(3, 2, "时段")
    for col, day in enumerate("一二三四五六日", start=3):
        sheet.cell(3, col, f"星期{day}")
    dayparts = ["上午", "上午", "下午", "下午", "晚上", "晚上"]
    numerals = ["一", "二", "三", "四", "五", "六"]
    for i, cell_text in enumerate(cell_texts):
        row = 4 + i
        sheet.cell(row, 1, numerals[i])
        sheet.cell(row, 2, dayparts[i])
        sheet.cell(row, 3, cell_text)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def _workbook_bytes() -> bytes:
    return _xlsx_with_courses(
        (
            "[1001]高等数学 [1-16周][1-2节] 教A101\n"
            "[FX3001004]财务管理Ⅰ（辅修） [1-16周][1-2节] 教C404"
        ),
        "[1002]程序设计 [1-16周][3-4节] 教B202",
    )


def test_ignore_minor_defaults_to_true_and_deselects_minor_rows(
    session: Session,
) -> None:
    service = ImportService(session)
    result = service.ingest("t.xlsx", _workbook_bytes())

    assert result["status"] in (ImportStatus.PARSED, ImportStatus.NEEDS_REVIEW)
    preview = service.preview(result["import_id"])
    assert preview["rows"]

    by_name = {row["course_name"]: row for row in preview["rows"]}
    assert "高等数学" in by_name
    assert "财务管理Ⅰ（辅修）" in by_name

    normal = by_name["高等数学"]
    minor = by_name["财务管理Ⅰ（辅修）"]
    assert normal["selected"] is True
    assert minor["selected"] is False

    minor_codes = {issue["code"] for issue in minor["issues"]}
    assert "MINOR_COURSE_EXCLUDED" in minor_codes
    excluded = [
        issue
        for issue in minor["issues"]
        if issue["code"] == "MINOR_COURSE_EXCLUDED"
    ]
    assert excluded and all(issue["severity"] == "info" for issue in excluded)
    assert not any(issue["code"] == "MINOR_COURSE_EXCLUDED" for issue in normal["issues"])
    # info issue must not flip row_status away from PARSED
    assert minor["status"] == "PARSED"

    assert preview.get("ignore_minor") is True


def test_ignore_minor_false_keeps_all_selected(session: Session) -> None:
    service = ImportService(session)
    result = service.ingest("t.xlsx", _workbook_bytes(), ignore_minor=False)

    preview = service.preview(result["import_id"])
    by_name = {row["course_name"]: row for row in preview["rows"]}
    assert by_name["财务管理Ⅰ（辅修）"]["selected"] is True
    assert not any(
        issue["code"] == "MINOR_COURSE_EXCLUDED"
        for issue in by_name["财务管理Ⅰ（辅修）"]["issues"]
    )
    assert preview.get("ignore_minor") is False


def test_confirm_skips_deselected_minor_and_writes_normal(
    session: Session,
) -> None:
    service = ImportService(session)
    result = service.ingest("t.xlsx", _workbook_bytes())
    confirmed = service.confirm(result["import_id"], class_name="计科2501")

    assert confirmed["status"] == ImportStatus.CONFIRMED
    # only 高等数学 + 程序设计 (minor deselected)
    assert confirmed["rows_confirmed"] == 2
    names = {row.name for row in session.query(CourseRow).all()}
    assert names == {"高等数学", "程序设计"}
    assert session.query(MeetingOccurrenceRow).count() == 2


def test_explicit_row_ids_can_reselect_minor_course(session: Session) -> None:
    """审核页改回勾选（显式 row_ids）后，曾被 ignore_minor 取消的行应能入库。"""
    service = ImportService(session)
    result = service.ingest("t.xlsx", _workbook_bytes())
    preview = service.preview(result["import_id"])
    minor_id = next(
        row["row_id"]
        for row in preview["rows"]
        if row["course_name"] == "财务管理Ⅰ（辅修）"
    )
    normal_ids = [
        row["row_id"]
        for row in preview["rows"]
        if row["course_name"] != "财务管理Ⅰ（辅修）"
    ]

    confirmed = service.confirm(
        result["import_id"],
        [*normal_ids, minor_id],
        class_name="计科2501",
    )

    assert confirmed["status"] == ImportStatus.CONFIRMED
    assert confirmed["rows_confirmed"] == 3
    names = {row.name for row in session.query(CourseRow).all()}
    assert "财务管理Ⅰ（辅修）" in names
